from smb.SMBConnection import SMBConnection
import tempfile
import PyPDF2
import docx
from openpyxl import load_workbook
from termcolor import colored
from pptx import Presentation
import src.modules.config as config


def conn_to_smb(user, pwd, server, ip, domain):
    client_machine_name = 'client' # as well any random name
    server_name = server
    conn = SMBConnection(user, pwd, client_machine_name, server_name, domain, use_ntlm_v2=True,
                        is_direct_tcp=True)
    conn.connect(ip, 445)
    shares = conn.listShares()
    return conn, shares

def list_files(conn, share_name, iter, authors_total, cnt):
    try:
        files = conn.listPath(share_name, iter)
        for f in files:
            if f.isDirectory:
                if f.filename not in [u'.', u'..']:
                    new_iter = iter + '{}\\'.format(f.filename)
                    list_files(conn, share_name, new_iter, authors_total, cnt) # recursion
            else:
                path = iter + f.filename
                if path.endswith(('.pdf', '.docx', '.xlsx', '.pptx')): 
                    cnt = retrieve_metadata(conn, share_name, path, authors_total, cnt)
    except Exception as e:
        return

def retrieve_metadata(conn, share_name, filepath, authors_total, cnt):
    authors = []
    if filepath.endswith('.xlsx'):
        file_obj = tempfile.NamedTemporaryFile(suffix='.xlsx')
    else:
        file_obj = tempfile.NamedTemporaryFile()                           
    conn.retrieveFile(share_name, filepath, file_obj)
    get_metadata(file_obj, filepath, authors)
    if len(authors) > 0:
        for author in authors:
            if cnt < 50:
                print(f"[+] {author}")
            authors_total.append(author)
            cnt+=1
    file_obj.close()
    return cnt

def get_metadata(file_obj, path, authors):
    if path.endswith('.pdf'):
        get_pdf_meta_author(file_obj.name, authors)
    elif path.endswith('.docx'):
        get_docx_meta_author(file_obj.name, authors)
    elif path.endswith('.xlsx'):
        get_xlsx_meta_author(file_obj.name, authors)
    elif path.endswith('.pptx'):
        get_pptx_meta_author(file_obj.name, authors)

def get_pdf_meta_author(path, authors):
    try:
        with open(path, 'rb') as pdf_file:
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            authors.append(pdf_reader.getDocumentInfo().author)
    except Exception as e:
        return

def get_docx_meta_author(path, authors):
    try:
        with open(path, 'rb'):
            doc = docx.Document(path)
            prop = doc.core_properties
            authors.append(prop.author)
            authors.append(prop.last_modified_by)
    except Exception as e:
        print(colored(f"[-] Error occurred while reading a Word file: {e}",'yellow'))
        return

def get_xlsx_meta_author(path, authors):
    try:
        with open(path, 'rb'):
            wb = load_workbook(path)
            prop = wb.properties
            authors.append(prop.creator)
            authors.append(prop.lastModifiedBy)
    except Exception as e:
        print(colored(f"[-] Error occurred while reading an Excel file: {e}",'yellow'))
        return

def get_pptx_meta_author(path, authors):
    try:
        with open(path, 'rb'):
            prs = Presentation(path)
            prop = prs.core_properties
            authors.append(prop.author)
            authors.append(prop.last_modified_by)
    except Exception as e:
        print(colored(f"[-] Error occurred while reading a Powerpoint file: {e}",'yellow'))
        return

def save_authors_to_file(authors_total):
    try:
        authors_total = list(set(authors_total))
        if not authors_total:
            return False
        with open("smb_meta_users.txt", "w") as f:
            for author in authors_total:
                f.write(author + '\n')
        return True
    except Exception as e:
        print(colored(f"[-] Error occurred: {e}",'yellow'))
        return False

def parse_users(input_file, output_file):
    try:
        with open(input_file, 'r') as infile, open(output_file, 'a') as outfile:
            for line in infile:
                names = line.strip().split()
                name = names[0][0]
                surname = names[1]
                samaccountname = (name + surname).lower()
                outfile.write(samaccountname + '\n')
        print(colored(f"[*] Users parsed to samaccountname form saved to {config.out_folder}/smb_meta_parsed_users.txt.",'green'))
    except Exception as e:
        print(colored(f"[-] Usernames could not be parsed to samaccountname form.",'yellow'))

def main(user, pwd, server, ip, domain):
    print(colored("\n[*] Trying to retrieve usernames from files' metadata.", 'blue'))
    authors_res = []
    conn, shares = conn_to_smb(user, pwd, server, ip, domain)
    cnt = 0
    for share in shares:
        if not share.isSpecial and share.name not in ['NETLOGON', 'SYSVOL', 'CertEnroll']: 
            print(colored(f"[*] Testing {share.name}...", 'blue'))
            authors = list_files(conn, share.name, '\\', authors_res, cnt)
    if save_authors_to_file(authors_res):
        print(colored(f"[+] Success! Unique usernames saved to {config.out_folder}/smb_meta_users.txt.", 'green'))
    else:
        print(colored(f"[-] No usernames retrieved.", 'yellow'))
    conn.close()
    # users to login format
    parse_users("smb_meta_users.txt", "smb_meta_parsed_users.txt")
